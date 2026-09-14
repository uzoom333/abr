import csv
import os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font


def exportar_excel():
    """Lê o historico.csv e exporta para um arquivo .xlsx com timestamp."""
    
    arquivo_csv = "historico.csv"
    
    # 1. Verifica se o CSV existe
    if not os.path.exists(arquivo_csv):
        print("Nenhuma consulta no historico ainda")
        return
    
    # 2. Cria a planilha
    wb = Workbook()                    # ← cria o Workbook
    ws = wb.active()      # ← pega a aba ativa
    
    # 3. Lê o CSV e escreve na planilha
    with open(arquivo_csv, "r") as f:
        reader = csv.reader(f)
        for linha in reader:
            ws.append(lista)         # ← adiciona a linha na planilha
    
    # 4. Deixa a primeira linha (cabeçalho) em negrito
    for celula in ws[1]:         # ws[1] pega TODAS as células da linha 1
        celula.font = Font(Bold = True)        # ← aplica negrito
    
    # 5. Gera nome do arquivo com timestamp
    agora = datetimenow           # ← pega data/hora atual formatada
    nome_arquivo = f"historico_{agora}.xlsx"
    
    # 6. Salva
    wb.save(???)                 # ← qual variável tem o nome?
    
    print(f"✅ Exportado para {nome_arquivo}")